import openpyxl as op

raw = op.load_workbook("COMM_BOM.xlsx")
wsRaw = raw.active

row_lists = []

for row in wsRaw.iter_rows(values_only=True):
    row_items = []

    for cell_value in row:
        if isinstance(cell_value, str):
            row_items.extend(
                x.strip() for x in cell_value.split(";") if x.strip()
            )

    row_lists.append(row_items)

print(row_lists)
row_lists[0].pop()
filtered_rows = []

for row in row_lists:
    if len(row) <= 4:
        # If 4 or fewer elements, keep the whole row
        filtered_rows.append(row)
    else:
        filtered_rows.append(row[:1] + row[-3:])

bom = op.Workbook()
ws = bom.active

for row_idx, values in enumerate(filtered_rows, start=1):
    for col_idx, value in enumerate(values, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)


bom.save("COMM_BOM.xlsx")